"""项目公共工具函数。"""

from __future__ import annotations

import importlib.util
import math
import os
import platform
import subprocess
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Callable

import numpy as np
import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows
from pypinyin import Style, lazy_pinyin, pinyin


CSV_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "gbk")
FILTER_OPERATORS = [
    "等于",
    "不等于",
    "包含",
    "不包含",
    "大于",
    "大于等于",
    "小于",
    "小于等于",
    "为空",
    "不为空",
]
NUMERIC_FILTER_OPERATORS = {"大于", "大于等于", "小于", "小于等于"}
BASE_STATS = [
    "数量",
    "缺失数",
    "缺失率",
    "平均值",
    "中位值",
    "最小值",
    "下四分位数",
    "上四分位数",
    "最大值",
    "极差",
    "标准差",
    "方差",
    "变异系数",
    "偏度",
    "峰度",
    "合计",
    "数量占比",
]
EXTENDED_STATS = BASE_STATS.copy()
COUNT_MASK_STATS = {"最大值", "最小值", "中位值", "下四分位数", "上四分位数", "极差"}
STD_MASK_STATS = {"标准差", "方差", "变异系数"}
SHAPE_MASK_STATS = {"偏度", "峰度"}
AREA_UNIT_TO_SQM = {
    "平方米": 1.0,
    "亩": 666.6666666667,
    "公顷": 10000.0,
    "平方千米": 1000000.0,
}
AREA_UNITS = list(AREA_UNIT_TO_SQM.keys())
DATE_GROUPINGS = ["年", "季度", "月", "周", "日"]


PluginCallable = Callable[[pd.Series], Any]


def sort_key(value: Any) -> tuple[str, int | float]:
    """生成适合中文排序的 key（拼音 + 笔画）。"""
    text = str(value).strip()
    if not text:
        return ("", float("inf"))

    first_char = text[0]
    pinyin_value = lazy_pinyin(first_char)[0]
    try:
        strokes = int(pinyin(first_char, style=Style.STROKES)[0][0])
    except Exception:
        strokes = float("inf")
    return (pinyin_value, strokes)



def round_to(value: Any, digits: int = 2) -> Any:
    try:
        if pd.isna(value):
            return value
        quant = "0" if digits <= 0 else "0." + "0" * digits
        return float(Decimal(str(value)).quantize(Decimal(quant), ROUND_HALF_UP))
    except Exception:
        return value



def round2(value: Any) -> Any:
    return round_to(value, 2)



def open_with_default_app(filepath: str) -> None:
    """使用系统默认程序打开文件。"""
    system = platform.system()
    if system == "Windows":
        os.startfile(filepath)  # type: ignore[attr-defined]
    elif system == "Darwin":
        subprocess.call(["open", filepath])
    else:
        subprocess.call(["xdg-open", filepath])



def read_csv_safely(path: str, **kwargs: Any) -> pd.DataFrame:
    """按常见中文环境编码依次尝试读取 CSV。"""
    errors: list[str] = []
    for encoding in CSV_ENCODINGS:
        try:
            return pd.read_csv(path, encoding=encoding, **kwargs)
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc}")
        except Exception:
            raise

    detail = "\n".join(errors) if errors else "未捕获到具体编码错误。"
    raise ValueError(f"无法识别 CSV 编码，可尝试另存为 UTF-8 后再导入。\n{detail}")



def friendly_error_message(exc: Exception) -> str:
    text = str(exc).strip() or exc.__class__.__name__
    low = text.lower()
    if "keyerror" in low or "not in index" in low:
        return f"字段不存在或字段名为空，请检查值字段、分组字段、面积字段是否已经正确选择。\n\n原始信息：{text}"
    if "No module named" in text:
        return f"缺少运行模块：{text}。如果是打包版本，请重新按新的 main.spec 打包。"
    if "At least one sheet must be visible" in text:
        return "导出过程中没有生成任何有效工作表，通常是因为没有可统计的字段或全部组合都被跳过。请检查值字段、面积字段和筛选条件。"
    if "Permission denied" in text:
        return f"文件正被其他程序占用或没有写入权限。\n\n原始信息：{text}"
    return text



def load_stat_plugins(plugin_dir: str | os.PathLike[str] = "plugins") -> dict[str, PluginCallable]:
    plugins: dict[str, PluginCallable] = {}
    path = Path(plugin_dir)
    if not path.exists() or not path.is_dir():
        return plugins

    for file in sorted(path.glob("*.py")):
        if file.name.startswith("_"):
            continue
        try:
            spec = importlib.util.spec_from_file_location(f"stat_plugin_{file.stem}", file)
            if spec is None or spec.loader is None:
                continue
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        except Exception:
            continue

        if hasattr(module, "register_stats"):
            try:
                mapping = module.register_stats()
                if isinstance(mapping, dict):
                    for name, func in mapping.items():
                        if callable(func):
                            plugins[str(name)] = func
            except Exception:
                continue
        elif hasattr(module, "STAT_NAME") and hasattr(module, "compute") and callable(module.compute):
            plugins[str(module.STAT_NAME)] = module.compute
    return plugins



def get_extended_stats_with_plugins(plugin_stats: dict[str, PluginCallable] | None = None) -> list[str]:
    names = BASE_STATS.copy()
    if plugin_stats:
        for name in plugin_stats:
            if name not in names:
                names.append(name)
    return names



def calculate_series_stats(
    series: pd.Series,
    total_valid_count: int | None = None,
    plugin_stats: dict[str, PluginCallable] | None = None,
) -> dict[str, Any]:
    """基于数值序列计算扩展统计指标。"""
    numeric = pd.to_numeric(series, errors="coerce")
    valid_count = int(numeric.count())
    total_count = int(len(numeric))
    missing_count = total_count - valid_count

    if valid_count:
        q25 = numeric.quantile(0.25)
        q75 = numeric.quantile(0.75)
        minimum = numeric.min()
        maximum = numeric.max()
        std = numeric.std()
        mean = numeric.mean()
        variance = numeric.var()
        total_sum = numeric.sum()
        skew = numeric.skew()
        kurt = numeric.kurt()
    else:
        q25 = np.nan
        q75 = np.nan
        minimum = np.nan
        maximum = np.nan
        std = np.nan
        mean = np.nan
        variance = np.nan
        total_sum = np.nan
        skew = np.nan
        kurt = np.nan

    count_pct = (valid_count / total_valid_count * 100) if total_valid_count else np.nan
    missing_rate = (missing_count / total_count * 100) if total_count else np.nan
    value_range = maximum - minimum if valid_count else np.nan
    cv = (std / mean) if valid_count and mean not in (0, 0.0) else np.nan

    result: dict[str, Any] = {
        "数量": valid_count,
        "缺失数": missing_count,
        "缺失率": round2(missing_rate),
        "平均值": round2(mean),
        "中位值": round2(numeric.median()) if valid_count else np.nan,
        "最小值": round2(minimum),
        "下四分位数": round2(q25),
        "上四分位数": round2(q75),
        "最大值": round2(maximum),
        "极差": round2(value_range),
        "标准差": round2(std),
        "方差": round2(variance),
        "变异系数": round2(cv),
        "偏度": round2(skew),
        "峰度": round2(kurt),
        "合计": round2(total_sum),
        "数量占比": round2(count_pct),
    }

    if plugin_stats:
        for name, func in plugin_stats.items():
            try:
                value = func(numeric.copy())
                if isinstance(value, (int, float, np.number)) and not isinstance(value, bool):
                    result[name] = round2(value)
                else:
                    result[name] = value
            except Exception:
                result[name] = np.nan

    return result



def apply_count_masks(df: pd.DataFrame) -> pd.DataFrame:
    """按样本量对不稳定指标做空值屏蔽。"""
    if "数量" not in df.columns:
        return df

    result = df.copy()
    counts = pd.to_numeric(result["数量"], errors="coerce").fillna(0)

    for name in COUNT_MASK_STATS:
        if name in result.columns:
            result.loc[counts <= 1, name] = np.nan

    for name in STD_MASK_STATS:
        if name in result.columns:
            result.loc[counts <= 5, name] = np.nan

    for name in SHAPE_MASK_STATS:
        if name in result.columns:
            result.loc[counts <= 7, name] = np.nan

    return result



def build_grouped_stats_frame(
    df: pd.DataFrame,
    field: str,
    group_fields: list[str],
    plugin_stats: dict[str, PluginCallable] | None = None,
) -> pd.DataFrame:
    """把 DataFrame 中某个值字段按分组计算成统计结果表。"""
    data = df.copy()
    data[field] = pd.to_numeric(data[field], errors="coerce")
    total_valid_count = int(data[field].count())

    if group_fields:
        rows: list[dict[str, Any]] = []
        for keys, group in data.groupby(group_fields, dropna=False):
            key_tuple = keys if isinstance(keys, tuple) else (keys,)
            row = {column: value for column, value in zip(group_fields, key_tuple)}
            row.update(calculate_series_stats(group[field], total_valid_count, plugin_stats))
            rows.append(row)
        return pd.DataFrame(rows)

    return pd.DataFrame([calculate_series_stats(data[field], total_valid_count, plugin_stats)])



def describe_filter_conditions(conditions: list[dict[str, Any]]) -> str:
    """把筛选条件列表转成便于显示的摘要。"""
    parts = []
    for condition in conditions:
        if not condition.get("field") or not condition.get("op"):
            continue
        field = str(condition["field"])
        op = str(condition["op"])
        value = str(condition.get("value", "")).strip()
        if op in {"为空", "不为空"}:
            parts.append(f"{field} {op}")
        else:
            parts.append(f"{field} {op} {value}")
    return "；".join(parts) if parts else "未启用"



def apply_filter_conditions(df: pd.DataFrame, conditions: list[dict[str, Any]]) -> pd.DataFrame:
    """按条件列表过滤 DataFrame。"""
    filtered = df.copy()
    for condition in conditions:
        field = str(condition.get("field", "")).strip()
        op = str(condition.get("op", "")).strip()
        value = str(condition.get("value", "")).strip()
        enabled = condition.get("enabled", True)

        if not enabled or not field or not op:
            continue
        if field not in filtered.columns:
            raise KeyError(f"筛选字段不存在：{field}")

        series = filtered[field]
        text_series = series.fillna("").astype(str).str.strip()

        if op == "为空":
            mask = series.isna() | (text_series == "")
        elif op == "不为空":
            mask = ~(series.isna() | (text_series == ""))
        elif op == "包含":
            mask = text_series.str.contains(value, case=False, na=False, regex=False)
        elif op == "不包含":
            mask = ~text_series.str.contains(value, case=False, na=False, regex=False)
        elif op in NUMERIC_FILTER_OPERATORS:
            try:
                target = float(value)
            except ValueError as exc:
                raise ValueError(f"筛选条件“{field} {op} {value}”不是有效数字。") from exc
            numeric_series = pd.to_numeric(series, errors="coerce")
            if op == "大于":
                mask = numeric_series > target
            elif op == "大于等于":
                mask = numeric_series >= target
            elif op == "小于":
                mask = numeric_series < target
            else:
                mask = numeric_series <= target
        else:
            numeric_series = pd.to_numeric(series, errors="coerce")
            numeric_target = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
            if not pd.isna(numeric_target) and numeric_series.notna().any():
                if op == "等于":
                    mask = numeric_series == numeric_target
                else:
                    mask = numeric_series != numeric_target
            else:
                if op == "等于":
                    mask = text_series == value
                else:
                    mask = text_series != value
        filtered = filtered[mask].copy()
    return filtered



def recommend_fields(df: pd.DataFrame) -> dict[str, list[str] | str]:
    recommendations: dict[str, list[str] | str] = {
        "numeric_fields": [],
        "date_fields": [],
        "category_fields": [],
        "recommended_value": "",
        "recommended_area": "",
        "recommended_date": "",
    }
    if df.empty:
        return recommendations

    numeric_candidates: list[str] = []
    date_candidates: list[str] = []
    category_candidates: list[str] = []

    for column in df.columns:
        series = df[column]
        non_null = series.dropna()
        if non_null.empty:
            continue

        numeric_rate = pd.to_numeric(series, errors="coerce").notna().mean()
        date_rate = pd.to_datetime(series, errors="coerce").notna().mean()
        nunique = int(series.nunique(dropna=True))
        ratio_unique = nunique / max(len(non_null), 1)

        if numeric_rate >= 0.85:
            numeric_candidates.append(column)
        if date_rate >= 0.85:
            date_candidates.append(column)
        if ratio_unique <= 0.5 or nunique <= 50:
            category_candidates.append(column)

    recommendations["numeric_fields"] = numeric_candidates
    recommendations["date_fields"] = date_candidates
    recommendations["category_fields"] = category_candidates

    if numeric_candidates:
        area_like = [c for c in numeric_candidates if any(k in str(c).lower() for k in ["面积", "area", "亩", "sqm", "ha"])]
        recommendations["recommended_area"] = area_like[0] if area_like else numeric_candidates[0]
        non_area = [c for c in numeric_candidates if c != recommendations["recommended_area"]]
        recommendations["recommended_value"] = (non_area[0] if non_area else numeric_candidates[0])
    if date_candidates:
        recommendations["recommended_date"] = date_candidates[0]
    return recommendations



def apply_mapping_rules(df: pd.DataFrame, mapping_rules: list[dict[str, Any]]) -> pd.DataFrame:
    if not mapping_rules:
        return df
    result = df.copy()
    for rule in mapping_rules:
        if not isinstance(rule, dict) or not rule.get("enabled", True):
            continue
        source = str(rule.get("source", "")).strip()
        output = str(rule.get("output", "")).strip()
        rule_type = str(rule.get("type", "exact")).strip()
        if not source or source not in result.columns or not output:
            continue
        series = result[source]
        if rule_type == "exact":
            mapping = {str(k): v for k, v in dict(rule.get("mapping", {})).items()}
            result[output] = series.map(lambda x: mapping.get(str(x), x))
        elif rule_type == "bins":
            breaks = rule.get("breaks", [])
            try:
                values = [float(v) for v in breaks]
                if len(values) < 2:
                    continue
            except Exception:
                continue
            labels = rule.get("labels", []) or [f"{values[i]}~{values[i+1]}" for i in range(len(values)-1)]
            numeric = pd.to_numeric(series, errors="coerce")
            result[output] = pd.cut(numeric, bins=values, labels=labels[:len(values)-1], include_lowest=True, right=False)
        elif rule_type == "date":
            granularity = str(rule.get("granularity", "月"))
            result[output] = build_date_group_series(series, granularity)
    return result



def build_date_group_series(series: pd.Series, granularity: str) -> pd.Series:
    dt = pd.to_datetime(series, errors="coerce")
    if granularity == "年":
        return dt.dt.strftime("%Y")
    if granularity == "季度":
        return dt.dt.to_period("Q").astype(str)
    if granularity == "月":
        return dt.dt.strftime("%Y-%m")
    if granularity == "周":
        return dt.dt.to_period("W").astype(str)
    return dt.dt.strftime("%Y-%m-%d")



def get_area_multiplier(source_unit: str, target_unit: str) -> float:
    source = AREA_UNIT_TO_SQM.get(source_unit, AREA_UNIT_TO_SQM["平方米"])
    target = AREA_UNIT_TO_SQM.get(target_unit, AREA_UNIT_TO_SQM["平方米"])
    return source / target



def convert_area_series(series: pd.Series, source_unit: str, target_unit: str) -> pd.Series:
    multiplier = get_area_multiplier(source_unit, target_unit)
    numeric = pd.to_numeric(series, errors="coerce")
    return numeric * multiplier



def area_column_name(target_unit: str) -> str:
    return f"面积（{target_unit}）"



def merge_data_files(paths: list[str], sheet_name: str | None = None, include_source: bool = True) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    frames: list[pd.DataFrame] = []
    summary: list[dict[str, Any]] = []
    for path in paths:
        ext = os.path.splitext(path)[1].lower()
        if ext == ".csv":
            df = read_csv_safely(path)
            used_sheet = "CSV"
        else:
            excel = pd.ExcelFile(path)
            chosen = sheet_name if sheet_name in excel.sheet_names else excel.sheet_names[0]
            df = pd.read_excel(path, sheet_name=chosen)
            used_sheet = chosen
        if include_source:
            df["来源文件"] = os.path.basename(path)
        frames.append(df)
        summary.append({"文件": os.path.basename(path), "工作表": used_sheet, "记录数": len(df), "字段数": len(df.columns)})
    if not frames:
        return pd.DataFrame(), []
    merged = pd.concat(frames, ignore_index=True, sort=False)
    return merged, summary



def add_chart_sheet_from_frame(
    workbook,
    title: str,
    data: pd.DataFrame,
    category_columns: list[str],
    metric_candidates: list[str],
) -> str | None:
    if data.empty:
        return None
    categories = [c for c in category_columns if c in data.columns]
    metric = next((m for m in metric_candidates if m in data.columns), None)
    if metric is None:
        numeric_cols = [c for c in data.columns if c not in categories and pd.api.types.is_numeric_dtype(data[c])]
        metric = numeric_cols[0] if numeric_cols else None
    if metric is None:
        return None

    chart_df = data.copy()
    if categories:
        chart_df["分类"] = chart_df[categories].astype(str).agg(" / ".join, axis=1)
    else:
        chart_df["分类"] = [metric] * len(chart_df)
    chart_df = chart_df[["分类", metric]].copy()
    chart_df = chart_df.head(20)

    safe_title = f"{title}_图表"[:31]
    if safe_title in workbook.sheetnames:
        del workbook[safe_title]
    ws = workbook.create_sheet(safe_title)
    for row in dataframe_to_rows(chart_df, index=False, header=True):
        ws.append(row)

    data_ref = Reference(ws, min_col=2, min_row=1, max_row=len(chart_df) + 1)
    cat_ref = Reference(ws, min_col=1, min_row=2, max_row=len(chart_df) + 1)

    bar = BarChart()
    bar.title = f"{metric} 柱状图"
    bar.y_axis.title = metric
    bar.x_axis.title = "分类"
    bar.add_data(data_ref, titles_from_data=True)
    bar.set_categories(cat_ref)
    bar.height = 8
    bar.width = 18
    ws.add_chart(bar, "D2")

    if len(chart_df) <= 10:
        pie = PieChart()
        pie.title = f"{metric} 饼图"
        pie.add_data(data_ref, titles_from_data=True)
        pie.set_categories(cat_ref)
        pie.height = 8
        pie.width = 18
        ws.add_chart(pie, "D20")

    return safe_title



def build_publish_readme(metadata: dict[str, Any]) -> str:
    lines = ["导出发布说明", "=" * 24]
    for key, value in metadata.items():
        lines.append(f"{key}: {value}")
    return "\n".join(lines) + "\n"
